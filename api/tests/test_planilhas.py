"""Import e export de planilha."""

import io

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from fitradeoff_api import create_app
from fitradeoff_api import spreadsheet
from fitradeoff_api.config import Settings

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

LINHAS_OK = [
    ["Fornecedores", "Qualidade", "Custo", "Prazo"],
    ["direcao", "max", "min", "min"],
    ["unidade", "", "R$/un", "dias"],
    ["Fornecedor A", 0.82, 12.5, 30],
    ["Fornecedor B", 0.95, 15.8, 12],
    ["Fornecedor C", 0.71, 10.2, 45],
]


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setenv("FITRADEOFF_DATA", str(tmp_path))
    from fitradeoff_api import deps

    deps.get_settings.cache_clear()
    app = create_app(Settings(storage=tmp_path))
    app.dependency_overrides[deps.get_settings] = lambda: Settings(storage=tmp_path)
    return TestClient(app)


def como_xlsx(linhas) -> bytes:
    livro = Workbook()
    for linha in linhas:
        livro.active.append(linha)
    saida = io.BytesIO()
    livro.save(saida)
    return saida.getvalue()


def como_csv(linhas, separador=",") -> bytes:
    return "\n".join(
        separador.join("" if c is None else str(c) for c in linha) for linha in linhas
    ).encode("utf-8")


# --- leitura ------------------------------------------------------------------


def test_le_o_layout_esperado():
    p = spreadsheet.parse_linhas(LINHAS_OK)
    assert p.nome == "Fornecedores"
    assert p.alternativas == ["Fornecedor A", "Fornecedor B", "Fornecedor C"]
    assert [c["direcao"] for c in p.criterios] == ["max", "min", "min"]
    assert [c["unidade"] for c in p.criterios] == ["", "R$/un", "dias"]
    assert p.consequencias[0] == [0.82, 12.5, 30]
    assert p.custos is None


def test_chaves_sao_derivadas_do_nome_e_nao_colidem():
    p = spreadsheet.parse_linhas(
        [
            ["X", "Custo total", "Custo total", "Prazo"],
            ["direcao", "min", "min", "min"],
            ["A", 1, 2, 3],
            ["B", 4, 5, 6],
        ]
    )
    chaves = [c["chave"] for c in p.criterios]
    assert chaves == ["custo_total", "custo_total_2", "prazo"]


def test_aceita_virgula_decimal_do_excel_brasileiro():
    p = spreadsheet.parse_linhas(
        [
            ["X", "A", "B"],
            ["direcao", "max", "min"],
            ["Alt 1", "0,82", "12,5"],
            ["Alt 2", "0,95", "15,8"],
        ]
    )
    assert p.consequencias[0] == [0.82, 12.5]


def test_linha_de_custo_alimenta_o_portfolio():
    p = spreadsheet.parse_linhas(LINHAS_OK + [["custo do projeto", 10, 12, 8]])
    assert p.custos == [10.0, 12.0, 8.0]
    assert "custo do projeto" not in p.alternativas


def test_linhas_vazias_sao_ignoradas():
    com_buracos = [LINHAS_OK[0], [], LINHAS_OK[1], LINHAS_OK[2], [], *LINHAS_OK[3:]]
    assert len(spreadsheet.parse_linhas(com_buracos).alternativas) == 3


@pytest.mark.parametrize(
    "linhas,esperado",
    [
        ([["X", "A", "B"], ["direcao", "max", "min"]], "ao menos duas alternativas"),
        # falta so a linha direcao: a mensagem tem que dizer isso, nao
        # reclamar do numero de linhas
        (
            [["X", "A", "B"], ["p", 1, 2], ["q", 3, 4]],
            "Falta a linha 'direcao'",
        ),
        (
            [["X", "A", "B"], ["unidade", "", ""], ["p", 1, 2], ["q", 3, 4]],
            "Falta a linha 'direcao'",
        ),
        (
            [["X", "A", "B"], ["direcao", "maior", "min"], ["p", 1, 2], ["q", 3, 4]],
            "use 'max'",
        ),
        (
            [["X", "A", "B"], ["direcao", "max", "min"], ["p", 1, "abc"], ["q", 3, 4]],
            "nao e um numero",
        ),
        (
            [["X", "A", "B"], ["direcao", "max", "min"], ["p", 1], ["q", 3, 4]],
            "Falta o desempenho",
        ),
        (
            [["X", "A", "B"], ["direcao", "max", "min"], ["p", 1, 2], ["p", 3, 4]],
            "mais de uma vez",
        ),
        ([["X", "A"], ["direcao", "max"], ["p", 1], ["q", 2]], "dois criterios"),
    ],
)
def test_planilha_malformada_explica_o_problema(linhas, esperado):
    with pytest.raises(spreadsheet.PlanilhaInvalida, match=esperado):
        spreadsheet.parse_linhas(linhas)


def test_xlsx_e_csv_produzem_o_mesmo_problema():
    do_xlsx = spreadsheet.parse("x.xlsx", como_xlsx(LINHAS_OK))
    do_csv = spreadsheet.parse("x.csv", como_csv(LINHAS_OK))
    assert do_xlsx == do_csv


def test_csv_com_ponto_e_virgula():
    p = spreadsheet.parse("x.csv", como_csv(LINHAS_OK, separador=";"))
    assert len(p.alternativas) == 3


def test_xls_antigo_orienta_a_conversao():
    with pytest.raises(spreadsheet.PlanilhaInvalida, match="salve como .xlsx"):
        spreadsheet.parse("antigo.xls", b"qualquer coisa")


def test_extensao_desconhecida():
    with pytest.raises(spreadsheet.PlanilhaInvalida, match="nao reconhecida"):
        spreadsheet.parse("dados.pdf", b"x")


def test_arquivo_corrompido_nao_vaza_erro_tecnico():
    with pytest.raises(spreadsheet.PlanilhaInvalida, match="97-2003"):
        spreadsheet.parse("x.xlsx", b"isto nao e um zip")


# --- modelo -------------------------------------------------------------------


def test_modelo_e_legivel_pelo_proprio_leitor():
    """O que a plataforma entrega tem que voltar sem edicao."""
    p = spreadsheet.parse("modelo.xlsx", spreadsheet.modelo())
    assert len(p.alternativas) == 4
    assert [c["direcao"] for c in p.criterios] == ["max", "min", "min"]


def test_modelo_tem_guia_de_preenchimento():
    livro = load_workbook(io.BytesIO(spreadsheet.modelo()))
    assert "Como preencher" in livro.sheetnames


# --- rotas --------------------------------------------------------------------


def test_baixar_modelo(client):
    r = client.get("/api/modelo-planilha")
    assert r.status_code == 200
    assert r.headers["content-type"] == XLSX
    assert "modelo-fitradeoff.xlsx" in r.headers["content-disposition"]


def test_importar_cria_a_decisao(client):
    r = client.post(
        "/api/decisoes/importar",
        files={"arquivo": ("dados.xlsx", como_xlsx(LINHAS_OK), XLSX)},
    )
    assert r.status_code == 201
    corpo = r.json()
    assert corpo["nome"] == "Fornecedores"
    assert len(corpo["criterios"]) == 3
    assert corpo["criterios"][1]["unidade"] == "R$/un"


def test_importar_planilha_invalida_devolve_422_com_explicacao(client):
    ruim = [["X", "A", "B"], ["unidade", "", ""], ["p", 1, 2], ["q", 3, 4]]
    r = client.post(
        "/api/decisoes/importar",
        files={"arquivo": ("dados.csv", como_csv(ruim), "text/csv")},
    )
    assert r.status_code == 422
    assert "direcao" in r.json()["detail"]


def test_importar_arquivo_grande_demais_e_recusado(client):
    r = client.post(
        "/api/decisoes/importar",
        files={"arquivo": ("grande.csv", b"x" * (3 * 1024 * 1024), "text/csv")},
    )
    assert r.status_code == 413


def test_exportar_traz_entrada_resultado_e_auditoria(client):
    criada = client.post(
        "/api/decisoes/importar",
        files={"arquivo": ("dados.xlsx", como_xlsx(LINHAS_OK), XLSX)},
    ).json()
    pid = criada["id"]
    client.post(
        f"/api/decisoes/{pid}/ordenacao",
        json={"ordem": [c["chave"] for c in criada["criterios"]]},
    )
    client.post(f"/api/decisoes/{pid}/resposta", json={"preferencia": "A"})

    r = client.get(f"/api/decisoes/{pid}/exportar")
    assert r.status_code == 200
    livro = load_workbook(io.BytesIO(r.content))
    assert livro.sheetnames == ["Entrada", "Resultado", "Perguntas", "Auditoria"]

    trilha = livro["Auditoria"]
    declaracoes = [trilha.cell(l, 4).value for l in range(2, trilha.max_row + 1)]
    assert any(d and "Ordenou os critérios" in d for d in declaracoes)


def test_exportacao_volta_pelo_importador(client):
    """A aba Entrada do export precisa ser reimportavel."""
    criada = client.post(
        "/api/decisoes/importar",
        files={"arquivo": ("dados.xlsx", como_xlsx(LINHAS_OK), XLSX)},
    ).json()
    exportado = client.get(f"/api/decisoes/{criada['id']}/exportar").content

    voltou = spreadsheet.parse("export.xlsx", exportado)
    assert voltou.alternativas == ["Fornecedor A", "Fornecedor B", "Fornecedor C"]
    assert [c["direcao"] for c in voltou.criterios] == ["max", "min", "min"]


def test_exportar_decisao_inexistente(client):
    assert client.get("/api/decisoes/fantasma/exportar").status_code == 404


def test_problema_traz_a_matriz_de_valor_normalizada(client):
    """Os graficos comparam valor de 0 a 1, nao a unidade original: sem este
    campo eles renderizam vazio, e nenhum outro teste percebe."""
    r = client.post(
        "/api/decisoes",
        json={
            "nome": "t",
            "alternativas": ["A", "B", "C"],
            "criterios": [
                {"chave": "q", "nome": "Q", "direcao": "max"},
                {"chave": "c", "nome": "C", "direcao": "min"},
            ],
            "consequencias": [[0.0, 30.0], [10.0, 10.0], [5.0, 20.0]],
        },
    )
    assert r.status_code == 201
    corpo = r.json()
    assert "valores" in corpo, "a matriz normalizada nao chegou ao cliente"

    valores = corpo["valores"]
    assert set(valores) == {"A", "B", "C"}
    # maior e melhor: A e o pior, B o melhor
    assert valores["A"]["q"] == 0.0 and valores["B"]["q"] == 1.0
    # menor e melhor: a direcao inverte
    assert valores["A"]["c"] == 0.0 and valores["B"]["c"] == 1.0
    # o intermediario fica no meio
    assert valores["C"]["q"] == pytest.approx(0.5)
    for linha in valores.values():
        for v in linha.values():
            assert 0.0 <= v <= 1.0


def test_export_traz_o_resumo_tabular_das_perguntas(client):
    """A trilha narrada conta a decisao em prosa; esta aba e para conferir em
    qual ciclo cada resposta entrou."""
    criada = client.post(
        "/api/decisoes/importar",
        files={"arquivo": ("dados.xlsx", como_xlsx(LINHAS_OK), XLSX)},
    ).json()
    pid = criada["id"]
    client.post(
        f"/api/decisoes/{pid}/ordenacao",
        json={"ordem": [c["chave"] for c in criada["criterios"]]},
    )
    client.post(f"/api/decisoes/{pid}/resposta", json={"preferencia": "A"})

    livro = load_workbook(io.BytesIO(client.get(f"/api/decisoes/{pid}/exportar").content))
    aba = livro["Perguntas"]
    assert [c.value for c in aba[1]] == ["Ciclo", "Declaracao", "Autor", "Registrado em"]
    ciclos = [aba.cell(l, 1).value for l in range(2, aba.max_row + 1)]
    assert ciclos == [1, 2]


def test_export_com_sensibilidade_acrescenta_a_aba(client):
    criada = client.post(
        "/api/decisoes/importar",
        files={"arquivo": ("dados.xlsx", como_xlsx(LINHAS_OK), XLSX)},
    ).json()
    pid = criada["id"]
    client.post(
        f"/api/decisoes/{pid}/ordenacao",
        json={"ordem": [c["chave"] for c in criada["criterios"]]},
    )

    sem = load_workbook(io.BytesIO(client.get(f"/api/decisoes/{pid}/exportar").content))
    assert "Sensibilidade" not in sem.sheetnames

    com = load_workbook(
        io.BytesIO(client.get(f"/api/decisoes/{pid}/exportar?sensibilidade=true").content)
    )
    assert "Sensibilidade" in com.sheetnames
    aba = com["Sensibilidade"]
    rotulos = [aba.cell(l, 1).value for l in range(1, aba.max_row + 1)]
    assert "Estabilidade" in rotulos


def test_sensibilidade_no_export_e_ignorada_antes_de_comecar(client):
    """Sem elicitacao nao ha o que variar: o parametro nao pode quebrar."""
    criada = client.post(
        "/api/decisoes/importar",
        files={"arquivo": ("dados.xlsx", como_xlsx(LINHAS_OK), XLSX)},
    ).json()
    r = client.get(f"/api/decisoes/{criada['id']}/exportar?sensibilidade=true")
    assert r.status_code == 200
    assert "Sensibilidade" not in load_workbook(io.BytesIO(r.content)).sheetnames


# --- configuracao junto da planilha -------------------------------------------


def test_importar_sorting_aceita_os_perfis_do_formulario(client):
    """A planilha descreve o problema; a leitura que se quer dele vem da tela.

    Antes a importacao nao repassava `configuracao`, entao escolher 'sorting' e
    subir uma planilha falhava com 400 -- o motor exige perfis, e nao havia por
    onde informa-los.
    """
    r = client.post(
        "/api/decisoes/importar",
        files={"arquivo": ("dados.xlsx", como_xlsx(LINHAS_OK), XLSX)},
        data={
            "problematica": "sorting",
            "configuracao": '{"perfis": [0.35, 0.7], "classes": ["Baixa", "Alta", "Topo"]}',
        },
    )
    assert r.status_code == 201, r.json()
    assert r.json()["configuracao"]["perfis"] == [0.35, 0.7]


def test_importar_com_configuracao_ilegivel_explica(client):
    r = client.post(
        "/api/decisoes/importar",
        files={"arquivo": ("dados.xlsx", como_xlsx(LINHAS_OK), XLSX)},
        data={"problematica": "sorting", "configuracao": "{nao e json"},
    )
    assert r.status_code == 422
    assert "JSON" in r.json()["detail"]


def test_importar_sem_configuracao_segue_funcionando(client):
    """O caminho antigo nao pode ter mudado: escolha nao precisa de nada."""
    r = client.post(
        "/api/decisoes/importar",
        files={"arquivo": ("dados.xlsx", como_xlsx(LINHAS_OK), XLSX)},
    )
    assert r.status_code == 201


# --- previa: ler sem gravar ---------------------------------------------------


def test_previa_devolve_o_que_leu_sem_criar_nada(client):
    """Conferir antes de gravar so vale se nada for gravado na conferencia."""
    antes = len(client.get("/api/decisoes").json()["decisoes"])

    r = client.post(
        "/api/planilhas/previa",
        files={"arquivo": ("dados.xlsx", como_xlsx(LINHAS_OK), XLSX)},
    )
    assert r.status_code == 200
    corpo = r.json()
    assert corpo["alternativas"]
    assert corpo["criterios"][0]["direcao"] in ("max", "min")
    assert corpo["origem"] == "dados.xlsx"

    depois = len(client.get("/api/decisoes").json()["decisoes"])
    assert depois == antes, "a previa nao pode criar decisao"


def test_previa_aceita_texto_colado_do_excel(client):
    """A area de transferencia do Excel vem tabulada, nao com virgula."""
    colado = "\t".join(["Compras", "Custo", "Prazo"]) + "\n"
    colado += "\t".join(["direcao", "min", "min"]) + "\n"
    colado += "\t".join(["Fornecedor A", "120", "30"]) + "\n"
    colado += "\t".join(["Fornecedor B", "95", "45"]) + "\n"

    r = client.post("/api/planilhas/previa", data={"texto": colado})
    assert r.status_code == 200, r.json()
    corpo = r.json()
    assert corpo["nome"] == "Compras"
    assert corpo["alternativas"] == ["Fornecedor A", "Fornecedor B"]
    assert [c["direcao"] for c in corpo["criterios"]] == ["min", "min"]
    assert corpo["origem"] == "colado"


def test_previa_sem_arquivo_nem_texto_orienta(client):
    r = client.post("/api/planilhas/previa", data={"texto": "   "})
    assert r.status_code == 422
    assert "cole" in r.json()["detail"].lower()


def test_previa_com_planilha_invalida_explica_o_que_falta(client):
    ruim = [["X", "A", "B"], ["unidade", "", ""], ["p", 1, 2], ["q", 3, 4]]
    r = client.post(
        "/api/planilhas/previa",
        files={"arquivo": ("dados.csv", como_csv(ruim), "text/csv")},
    )
    assert r.status_code == 422
    assert "direcao" in r.json()["detail"]


def test_export_sob_veto_leva_a_ressalva_junto(client):
    """A planilha publica as faixas penalizadas; sem a ressalva ao lado, quem
    abrir o arquivo lê um teto conservador como se fosse exato.

    A ressalva chega aqui de graça porque o motor a publica sob a mesma chave
    que as carteiras já usavam -- é sempre uma afirmação sobre como o número
    foi obtido. Este teste existe para que essa unificação não se desfaça."""
    criada = client.post(
        "/api/decisoes",
        json={
            "nome": "Com veto",
            "alternativas": ["A", "B", "C"],
            "criterios": [
                {"chave": "q", "nome": "Qualidade", "direcao": "max"},
                {
                    "chave": "p",
                    "nome": "Prazo",
                    "direcao": "min",
                    "unidade": "dias",
                    "veto_superior": 30,
                    "veto_inferior": 50,
                },
            ],
            "consequencias": [[0.95, 60], [0.70, 20], [0.60, 25]],
        },
    ).json()
    client.post(f"/api/decisoes/{criada['id']}/ordenacao", json={"ordem": ["q", "p"]})

    r = client.get(f"/api/decisoes/{criada['id']}/exportar")
    assert r.status_code == 200
    aba = load_workbook(io.BytesIO(r.content))["Resultado"]
    linhas = {aba.cell(l, 1).value: aba.cell(l, 2).value for l in range(1, aba.max_row + 1)}

    assert "Ressalva" in linhas, "o export saiu com as faixas e sem dizer como saíram"
    assert "envelope" in linhas["Ressalva"]
